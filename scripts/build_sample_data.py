"""Generate the binary sample documents (PDF and DOCX) in data/sample/northwind.

The generated files are committed so evaluators need not run this; it exists so the
sample corpus is reproducible and reviewable as text. Run: python scripts/build_sample_data.py
"""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf

OUT = Path(__file__).resolve().parent.parent / "data" / "sample" / "northwind"

MANUAL_PAGES = [
    (
        "NW-200 Autonomous Warehouse Robot - Operator Manual",
        """This manual describes the safe operation and routine maintenance of the NW-200
autonomous mobile robot, manufactured by Northwind Robotics. Read it fully before
operating the robot. Keep it available to every operator on site.

The NW-200 moves shelving units and pallets inside warehouses. It navigates with a
360-degree lidar sensor and two depth cameras, and follows routes assigned by the
Northwind Fleet Manager software.

Technical specifications:
Maximum payload: 150 kg.
Maximum speed: 1.8 metres per second (limited to 1.0 m/s in zones marked for
pedestrians).
Battery: 48 V lithium iron phosphate, rated runtime of 8 hours per charge under typical
load.
Charging time: 90 minutes from 10% to 100% on the NW-C2 docking charger.
Operating temperature: 0 to 40 degrees Celsius.
Weight without payload: 210 kg.""",
    ),
    (
        "Safety",
        """Only operators who have completed the Northwind operator training may start,
stop or reassign an NW-200.

The red emergency stop buttons on each side of the robot cut power to the drive motors
immediately. After an emergency stop, the robot can only be restarted from the Fleet
Manager console, after an operator has confirmed that the area is clear.

Never ride on the robot, and never place a payload that blocks the front lidar sensor.
Keep a minimum clearance of 0.5 metres around a moving robot.

The robot must not be operated on ramps steeper than 5 degrees or on wet floors.""",
    ),
    (
        "Error codes",
        """When the status light turns amber or red, the Fleet Manager console shows an
error code. The most common codes are:

E-104 Lidar obstruction: the lidar field of view is blocked. Remove the obstruction and
clean the lidar window with a dry microfibre cloth. The robot resumes automatically.

E-221 Battery overheating: battery temperature above 55 degrees Celsius. The robot stops
and must cool down for at least 30 minutes. If E-221 occurs twice within one week,
contact field service; do not keep operating the robot.

E-310 Wheel motor stall: a drive wheel is blocked, often by shrink wrap or straps.
Press the emergency stop, remove the debris, then restart from the console.

E-502 Localisation lost: the robot cannot match its position to the site map. Push it
manually to the nearest marked docking point and restart it there.""",
    ),
    (
        "Maintenance schedule",
        """Daily: inspect the lidar window and depth cameras for dust, and check the wheels
for wrapped debris.

Every 500 operating hours: inspect the drive wheels for wear and replace them if the
tread depth is below 3 mm. Check the tightness of the payload lift bolts.

Every 2,000 operating hours: replace the lift actuator seals and run the battery health
diagnostic from the Fleet Manager.

Batteries should be replaced when their capacity falls below 70% of the original rating,
which is typically after about 3,000 charge cycles.

Warranty: the NW-200 has a 24-month warranty covering parts and labour. The warranty is
void if maintenance intervals are not followed or if non-Northwind parts are fitted.""",
    ),
]

RELEASE_NOTES = [
    ("heading", "NW-200 Firmware Release Notes"),
    ("para", "These notes cover firmware releases for the NW-200 warehouse robot."),
    ("heading", "Firmware 3.3 (June 2026)"),
    (
        "para",
        "Improved battery management: with firmware 3.3 the rated runtime per charge increases "
        "to 9.5 hours under typical load, thanks to smarter idle power saving.",
    ),
    (
        "para",
        "New error code E-615 Payload imbalance: raised when the load on the lift platform is "
        "off-centre by more than 15 cm. Re-centre the payload before continuing.",
    ),
    ("para", "Fixed an issue where the robot could report E-502 after a firmware update."),
    ("heading", "Firmware 3.2 (February 2026)"),
    (
        "para",
        "Added support for pedestrian zones: the maximum speed is automatically limited in zones "
        "marked as pedestrian areas in the Fleet Manager.",
    ),
    ("para", "Reduced charging time on the NW-C2 docking charger by 10 minutes."),
    (
        "table",
        [
            ("Version", "Release date", "Upgrade required"),
            ("3.3", "June 2026", "Recommended"),
            ("3.2", "February 2026", "Mandatory"),
        ],
    ),
]


def build_manual(path: Path) -> None:
    doc = pymupdf.open()
    for title, body in MANUAL_PAGES:
        page = doc.new_page()
        page.insert_text((56, 64), title, fontsize=15)
        rect = pymupdf.Rect(56, 88, page.rect.width - 56, page.rect.height - 56)
        page.insert_textbox(rect, body, fontsize=10.5, lineheight=1.35)
    doc.set_metadata({"title": "NW-200 Operator Manual", "author": "Northwind Robotics"})
    doc.save(path)


def build_release_notes(path: Path) -> None:
    document = docx.Document()
    for kind, content in RELEASE_NOTES:
        if kind == "heading":
            document.add_heading(content, level=1)
        elif kind == "para":
            document.add_paragraph(content)
        else:
            table = document.add_table(rows=0, cols=len(content[0]))
            for row in content:
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=True):
                    cell.text = value
    document.save(path)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_manual(OUT / "nw200_operator_manual.pdf")
    build_release_notes(OUT / "firmware_release_notes.docx")
    print(f"wrote sample documents to {OUT}")
