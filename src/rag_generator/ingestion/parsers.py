"""Concrete parsers for PDF, plain text / Markdown, DOCX, CSV, Excel and PowerPoint.

Parsers work on bytes (not paths) so uploads never need to touch disk.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Sequence
from datetime import date, datetime, time

from rag_generator.errors import CorruptDocumentError, EmptyDocumentError, OCRRequiredError
from rag_generator.models import ParsedDocument, Section

# A page with fewer extractable characters than this, but with images, is treated as scanned.
_SCANNED_PAGE_MAX_CHARS = 20
_WHITESPACE_RUNS = re.compile(r"[ \t]+")
_BLANK_LINE_RUNS = re.compile(r"\n{3,}")


def normalise_whitespace(text: str) -> str:
    """Collapse horizontal whitespace and excess blank lines; keep paragraph breaks."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = _WHITESPACE_RUNS.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_LINE_RUNS.sub("\n\n", text).strip()


def decode_text(data: bytes) -> tuple[str, list[str]]:
    """UTF-8 (BOM tolerated), falling back to latin-1 with a warning."""
    try:
        return data.decode("utf-8-sig"), []
    except UnicodeDecodeError:
        return data.decode("latin-1"), ["not valid UTF-8; decoded as latin-1"]


class TextParser:
    """UTF-8 text and Markdown, with a latin-1 fallback for legacy encodings."""

    extensions = ("txt", "md", "markdown")

    def parse(self, data: bytes, source: str) -> ParsedDocument:
        text, warnings = decode_text(data)
        text = normalise_whitespace(text)
        if not text:
            raise EmptyDocumentError(f"{source}: file contains no text")
        file_type = source.rsplit(".", 1)[-1].lower() if "." in source else "txt"
        return ParsedDocument(
            source=source, file_type=file_type, sections=[Section(text=text)], warnings=warnings
        )


class PdfParser:
    """PDF via PyMuPDF; one section per page so citations can name the page."""

    extensions = ("pdf",)

    def parse(self, data: bytes, source: str) -> ParsedDocument:
        import pymupdf

        try:
            doc = pymupdf.open(stream=data, filetype="pdf")
        except Exception as exc:  # PyMuPDF raises several unrelated types for bad input
            raise CorruptDocumentError(f"{source}: not a readable PDF ({exc})") from exc

        with doc:
            if doc.needs_pass:
                raise CorruptDocumentError(f"{source}: PDF is password-protected")
            sections, scanned_pages = self._extract_pages(doc)

        warnings = []
        if scanned_pages:
            warnings.append(
                f"{len(scanned_pages)} page(s) look scanned (no text layer) and were skipped: "
                f"{scanned_pages[:10]}"
            )
        if not sections:
            if scanned_pages:
                raise OCRRequiredError(
                    f"{source}: no extractable text; the PDF appears to be scanned images. "
                    "OCR is not supported."
                )
            raise EmptyDocumentError(f"{source}: PDF contains no text")
        return ParsedDocument(source=source, file_type="pdf", sections=sections, warnings=warnings)

    @staticmethod
    def _extract_pages(doc) -> tuple[list[Section], list[int]]:
        sections: list[Section] = []
        scanned_pages: list[int] = []
        for page_number, page in enumerate(doc, start=1):
            text = normalise_whitespace(page.get_text("text", sort=True))
            if len(text) >= _SCANNED_PAGE_MAX_CHARS:
                sections.append(Section(text=text, page=page_number))
            elif page.get_images():
                scanned_pages.append(page_number)
            elif text:
                sections.append(Section(text=text, page=page_number))
        return sections, scanned_pages


class DocxParser:
    """DOCX paragraphs followed by table rows (cells joined with ' | ')."""

    extensions = ("docx",)

    def parse(self, data: bytes, source: str) -> ParsedDocument:
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:  # python-docx raises zipfile/KeyError/lxml errors
            raise CorruptDocumentError(f"{source}: not a readable DOCX ({exc})") from exc

        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        text = normalise_whitespace("\n\n".join(blocks))
        if not text:
            raise EmptyDocumentError(f"{source}: DOCX contains no text")
        return ParsedDocument(source=source, file_type="docx", sections=[Section(text=text)])


def _cell_text(value: object) -> str:
    """Render one spreadsheet cell on a single line; '' for an empty cell."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat(" ")
    if isinstance(value, date | time):
        return value.isoformat()
    return " ".join(str(value).split())


def format_table_rows(rows: Iterable[Sequence[object]], prefix: str = "") -> list[str]:
    """Flatten a table into one self-describing line per row.

    The first non-empty row is the header. Each later row becomes
    ``"{prefix}Row 7: Region: West | Revenue: 1200"`` (the 1-based row number in the file;
    empty cells omitted), so a chunk cut from the middle of a long table still carries its
    column names. A table with only a header yields a single ``Columns:`` line.
    """
    header: list[str] | None = None
    lines: list[str] = []
    for row_number, row in enumerate(rows, start=1):
        cells = [_cell_text(v) for v in row]
        if not any(cells):
            continue
        if header is None:
            header = cells
            continue
        pairs = [
            f"{header[i] if i < len(header) and header[i] else f'Column {i + 1}'}: {cell}"
            for i, cell in enumerate(cells)
            if cell
        ]
        lines.append(f"{prefix}Row {row_number}: " + " | ".join(pairs))
    if header is not None and not lines:
        lines.append(f"{prefix}Columns: " + " | ".join(c for c in header if c))
    return lines


class CsvParser:
    """CSV / TSV; one line per row with the header names repeated (``format_table_rows``)."""

    extensions = ("csv", "tsv")

    def parse(self, data: bytes, source: str) -> ParsedDocument:
        text, warnings = decode_text(data)
        file_type = source.rsplit(".", 1)[-1].lower()
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = "\t" if file_type == "tsv" else ","
        try:
            lines = format_table_rows(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error as exc:
            raise CorruptDocumentError(f"{source}: not a readable CSV ({exc})") from exc
        if not lines:
            raise EmptyDocumentError(f"{source}: CSV contains no data")
        return ParsedDocument(
            source=source,
            file_type=file_type,
            sections=[Section(text="\n".join(lines))],
            warnings=warnings,
        )


class XlsxParser:
    """Excel workbooks via openpyxl; one section per non-empty worksheet.

    Formula cells use the value Excel last saved. A workbook written by a tool that never
    calculated it (e.g. openpyxl itself) has no saved values, so those cells read as empty.
    Legacy ``.xls`` is not supported.
    """

    extensions = ("xlsx", "xlsm")

    def parse(self, data: bytes, source: str) -> ParsedDocument:
        import openpyxl

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # zipfile/KeyError/InvalidFileException for bad input
            raise CorruptDocumentError(f"{source}: not a readable Excel workbook ({exc})") from exc

        sections: list[Section] = []
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                lines = format_table_rows(rows, prefix=f"[{sheet.title}] ")
                if lines:
                    sections.append(Section(text="\n".join(lines)))
        except Exception as exc:  # malformed sheet XML surfaces lazily in read-only mode
            raise CorruptDocumentError(f"{source}: unreadable worksheet ({exc})") from exc
        finally:
            workbook.close()

        if not sections:
            raise EmptyDocumentError(f"{source}: workbook contains no data")
        file_type = source.rsplit(".", 1)[-1].lower()
        return ParsedDocument(source=source, file_type=file_type, sections=sections)


class PptxParser:
    """PowerPoint via python-pptx; one section per slide, numbered like pages.

    Collects text boxes, tables (cells joined with ' | '), grouped shapes and speaker
    notes. Charts and pictures carry no extractable text. Legacy ``.ppt`` is not supported.
    """

    extensions = ("pptx",)

    def parse(self, data: bytes, source: str) -> ParsedDocument:
        import pptx

        try:
            presentation = pptx.Presentation(io.BytesIO(data))
        except Exception as exc:  # zipfile/KeyError/lxml errors for bad input
            raise CorruptDocumentError(f"{source}: not a readable PPTX ({exc})") from exc

        sections: list[Section] = []
        textless: list[int] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            blocks = [block for shape in slide.shapes for block in _shape_text(shape)]
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    blocks.append(f"Notes: {notes.text}")
            text = normalise_whitespace("\n\n".join(blocks))
            if text:
                sections.append(Section(text=text, page=slide_number))
            else:
                textless.append(slide_number)

        if not sections:
            raise EmptyDocumentError(f"{source}: presentation contains no text")
        warnings = []
        if textless:
            warnings.append(
                f"{len(textless)} slide(s) have no extractable text and were skipped: "
                f"{textless[:10]}"
            )
        return ParsedDocument(source=source, file_type="pptx", sections=sections, warnings=warnings)


def _shape_text(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [block for child in shape.shapes for block in _shape_text(child)]
    if shape.has_table:
        rows = []
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        return ["\n".join(rows)] if rows else []
    if shape.has_text_frame and shape.text_frame.text.strip():
        return [shape.text_frame.text]
    return []
