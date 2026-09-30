"""Shared fixtures. Every test runs offline: hashing embeddings and a scripted fake LLM."""

from __future__ import annotations

import io
import json
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest

from rag_generator.config import Settings
from rag_generator.errors import GenerationError
from rag_generator.generation import LLMResponse
from rag_generator.orchestration import RAGApplication

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DIR = ROOT / "data" / "sample" / "northwind"
ALT_SAMPLE_DIR = ROOT / "data" / "sample" / "gardening"


class FakeLLM:
    """Scripted LLMProvider. ``responder`` receives (system, user, schema) and returns a
    dict, or raises. Calls are recorded for assertions."""

    def __init__(self, responder: Callable[[str, str, dict], dict[str, Any]]) -> None:
        self.responder = responder
        self.calls: list[dict[str, Any]] = []

    @property
    def model_id(self) -> str:
        return "fake-llm"

    def generate_json(self, system: str, user: str, schema: dict[str, Any]) -> LLMResponse:
        self.calls.append({"system": system, "user": user, "schema": schema})
        return LLMResponse(
            data=self.responder(system, user, schema),
            model="fake-llm",
            input_tokens=100,
            output_tokens=20,
        )


def cite_first_source(system: str, user: str, schema: dict) -> dict:
    """Answers with the text of S1 and cites it, like a well-behaved model."""
    if "queries" in schema["properties"]:
        return {"queries": ["alternative phrasing"]}
    match = re.search(r'<source id="S1"[^>]*>\n(.*?)\n</source>', user, re.S)
    text = match.group(1)[:80] if match else ""
    return {"answerable": True, "answer": f"{text} [S1]", "citations": ["S1"], "reason": ""}


def always_unanswerable(system: str, user: str, schema: dict) -> dict:
    if "queries" in schema["properties"]:
        return {"queries": ["rewritten query one", "rewritten query two"]}
    return {"answerable": False, "answer": "", "citations": [], "reason": "not in sources"}


def failing_llm(system: str, user: str, schema: dict) -> dict:
    raise GenerationError("LLM rate limit exceeded after retries", retryable=True)


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(
        data_dir=tmp_path / "data",
        embedding_provider="hashing",
        llm_provider="none",
        min_relevance=0.0,
        chunk_size=300,
        chunk_overlap=50,
        _env_file=None,
    )


@pytest.fixture
def make_app(settings: Settings) -> Callable[..., RAGApplication]:
    def _make(llm: FakeLLM | None = None, **overrides: Any) -> RAGApplication:
        s = settings.model_copy(update=overrides) if overrides else settings
        return RAGApplication(s, llm=llm)

    return _make


# --- Document builders ---------------------------------------------------------------


def make_pdf(pages: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_image_only_pdf() -> bytes:
    """A PDF whose only page is an image (simulates a scanned document)."""
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 64, 64), False)
    pix.clear_with(200)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(50, 50, 300, 300), pixmap=pix)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=0, cols=len(table[0]))
        for row in table:
            for cell, value in zip(t.add_row().cells, row, strict=True):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list[Any]]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pptx(
    slides: list[list[str]],
    notes: dict[int, str] | None = None,
    table: list[list[str]] | None = None,
) -> bytes:
    """One slide per entry (each string a text box). ``notes`` maps 1-based slide number to
    speaker notes; ``table`` is added to the last slide. An empty entry is a blank slide."""
    import pptx
    from pptx.util import Inches

    presentation = pptx.Presentation()
    blank = presentation.slide_layouts[6]
    for number, texts in enumerate(slides, start=1):
        slide = presentation.slides.add_slide(blank)
        for i, text in enumerate(texts):
            box = slide.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(6), Inches(1))
            box.text_frame.text = text
        if notes and number in notes:
            slide.notes_slide.notes_text_frame.text = notes[number]
        if table and number == len(slides):
            shape = slide.shapes.add_table(
                len(table), len(table[0]), Inches(1), Inches(4), Inches(6), Inches(2)
            )
            for r, row in enumerate(table):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def dump_jsonl(path: Path, rows: list[dict]) -> Path:
    path.write_text("\n".join(json.dumps(r) for r in rows), encoding="utf-8")
    return path
